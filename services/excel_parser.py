import os
import json
import openpyxl
from config import FILTERABLE_THRESHOLD
from models.database import get_db


def parse_and_import(filepath, filename):
    """
    Parse an Excel file and import it into the database.
    Returns import_id, row_count, and column configs.
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    # Read column headers from first row
    headers = []
    for cell in ws[1]:
        val = str(cell.value).strip() if cell.value is not None else "col{}".format(cell.column)
        headers.append(val)

    # Read all data rows
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        row_dict = {}
        for i, val in enumerate(row):
            if i < len(headers):
                row_dict[headers[i]] = str(val).strip() if val is not None else ''
        rows.append(row_dict)

    # Count unique values per column
    col_values = {h: set() for h in headers}
    for row_dict in rows:
        for h in headers:
            v = row_dict.get(h, '')
            if v:
                col_values[h].add(v)

    conn = get_db()
    try:
        c = conn.cursor()

        # Insert import record
        c.execute(
            "INSERT INTO excel_imports (filename, filepath, row_count) VALUES (?, ?, ?)",
            (filename, filepath, len(rows))
        )
        import_id = c.lastrowid

        # Insert cards
        for idx, row_dict in enumerate(rows):
            c.execute(
                "INSERT INTO cards (import_id, row_index, data) VALUES (?, ?, ?)",
                (import_id, idx + 2, json.dumps(row_dict, ensure_ascii=False))
            )

        # Insert column_configs
        for order, col in enumerate(headers):
            unique_count = len(col_values[col])
            total_non_empty = sum(1 for r in rows if r.get(col, ''))
            # High-cardinality columns (>60% unique values) are likely free text, not filterable
            high_cardinality = (total_non_empty > 0 and unique_count / total_non_empty > 0.6)
            is_filterable = 1 if (0 < unique_count <= FILTERABLE_THRESHOLD and not high_cardinality) else 0
            field_type = 'tag' if is_filterable else 'text'
            is_primary = 1 if order == 0 else 0
            c.execute(
                "INSERT OR IGNORE INTO column_configs "
                "(import_id, column_name, display_name, is_visible, is_filterable, "
                "display_order, is_primary, field_type) "
                "VALUES (?, ?, ?, 1, ?, ?, ?, ?)",
                (import_id, col, col, is_filterable, order, is_primary, field_type)
            )

        # Insert filter_options (only for filterable columns)
        for col in headers:
            unique_count = len(col_values[col])
            total_non_empty = sum(1 for r in rows if r.get(col, ''))
            high_cardinality = (total_non_empty > 0 and unique_count / total_non_empty > 0.6)
            if 0 < unique_count <= FILTERABLE_THRESHOLD and not high_cardinality:
                for val in sorted(col_values[col]):
                    count = sum(1 for r in rows if r.get(col, '') == val)
                    c.execute(
                        "INSERT OR IGNORE INTO filter_options "
                        "(import_id, column_name, value, card_count) VALUES (?, ?, ?, ?)",
                        (import_id, col, val, count)
                    )

        # Update active import
        c.execute(
            "INSERT OR REPLACE INTO app_settings (key, value) VALUES ('active_import_id', ?)",
            (str(import_id),)
        )

        conn.commit()

        cols = c.execute(
            "SELECT * FROM column_configs WHERE import_id=? ORDER BY display_order",
            (import_id,)
        ).fetchall()

        return {
            'import_id': import_id,
            'row_count': len(rows),
            'columns': [dict(col) for col in cols]
        }
    finally:
        conn.close()
